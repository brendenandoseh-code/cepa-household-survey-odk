"""Generate the CEPA household survey XLSForm.

The workbook is a build artifact. This script is the source of truth so the form is
diffable in git — an .xlsx is an opaque zip and reviewing changes to one is guesswork.

Run:  python build_xlsform.py
Then: python -m pyxform.xls2xform cepa_household_survey.xlsx cepa_household_survey.xml
"""

from openpyxl import Workbook

OUT = "cepa_household_survey.xlsx"

# --- survey sheet -----------------------------------------------------------
# Columns follow the XLSForm spec: https://xlsform.org/en/
SURVEY = [
    ["type", "name", "label", "hint", "required", "relevant", "constraint",
     "constraint_message", "appearance", "default"],

    # -- metadata captured automatically by the client --
    ["start", "start", None, None, None, None, None, None, None, None],
    ["end", "end", None, None, None, None, None, None, None, None],
    ["today", "today", None, None, None, None, None, None, None, None],

    # -- enumeration context --
    ["begin_group", "visit", "Visit details", None, None, None, None, None, "field-list", None],
    ["text", "enumerator", "Enumerator name", "Person conducting this interview.",
     "yes", None, None, None, None, None],
    ["select_one village", "village", "Village", "Ha Monyake chiefdom covers three villages.",
     "yes", None, None, None, None, None],
    ["date", "visit_date", "Date of visit", None, "yes", None,
     ". <= today()", "The visit date cannot be in the future.", None, "today()"],
    ["end_group", "visit", None, None, None, None, None, None, None, None],

    # -- consent gate --
    ["note", "consent_intro",
     "Read aloud: I am collecting information about households in this village to help plan "
     "health and community projects. Your answers are confidential and no names are recorded. "
     "You may skip any question or stop at any time.",
     None, None, None, None, None, None, None],
    ["select_one yes_no", "consent", "Does the household representative agree to take part?",
     None, "yes", None, None, None, "minimal", None],
    ["note", "consent_declined", "No consent given. Please end the interview here and submit this form.",
     None, None, "${consent} = 'no'", None, None, None, None],

    # -- household composition --
    ["begin_group", "composition", "Household composition", None, None,
     "${consent} = 'yes'", None, None, None, None],
    ["integer", "household_size", "How many people live in this household?", None, "yes", None,
     ". >= 1 and . <= 30", "Household size must be between 1 and 30. Confirm the answer.",
     None, None],
    ["integer", "single_parents", "How many single parents live in this household?", None, "yes", None,
     ". >= 0 and . <= ${household_size}",
     "Cannot exceed the household size you entered.", None, None],
    ["integer", "people_over_65", "How many household members are aged 65 or over?", None, "yes", None,
     ". >= 0 and . <= ${household_size}",
     "Cannot exceed the household size you entered.", None, None],
    ["end_group", "composition", None, None, None, None, None, None, None, None],

    # -- sensitive block, deliberately optional --
    ["begin_group", "health", "Health", None, None, "${consent} = 'yes'", None, None, None, None],
    ["note", "hiv_note",
     "Read aloud: The next question is optional. You do not have to answer it.",
     None, None, None, None, None, None, None],
    ["integer", "people_with_hiv",
     "How many household members are living with HIV?",
     "Leave blank if the respondent prefers not to answer.",
     None, None, ". >= 0 and . <= ${household_size}",
     "Cannot exceed the household size you entered.", None, None],
    ["end_group", "health", None, None, None, None, None, None, None, None],

    # -- livelihoods and community assessment --
    ["begin_group", "livelihoods", "Livelihoods and community", None, None,
     "${consent} = 'yes'", None, None, None, None],
    ["select_multiple income_source", "income_source",
     "What are the main sources of household income?", "Select all that apply.",
     "yes", None, None, None, None, None],
    ["text", "income_source_other", "Please specify the other income source",
     None, "yes", "selected(${income_source}, 'other')", None, None, None, None],
    ["select_multiple strength", "strengths",
     "What skills or strengths does this household have?", "Select all that apply.",
     "yes", None, None, None, None, None],
    ["text", "strengths_other", "Please specify the other skill or strength",
     None, "yes", "selected(${strengths}, 'other')", None, None, None, None],
    ["select_multiple challenge", "challenges",
     "What are the main challenges facing this household?", "Select all that apply.",
     "yes", None, None, None, None, None],
    ["text", "challenges_other", "Please specify the other challenge",
     None, "yes", "selected(${challenges}, 'other')", None, None, None, None],
    ["select_multiple opportunity", "opportunities",
     "What would most help this community grow?", "Select all that apply.",
     "yes", None, None, None, None, None],
    ["text", "opportunities_other", "Please specify the other opportunity",
     None, "yes", "selected(${opportunities}, 'other')", None, None, None, None],
    ["text", "notes", "Interviewer notes", "Anything the categories above did not capture.",
     None, None, None, None, "multiline", None],
    ["end_group", "livelihoods", None, None, None, None, None, None, None, None],
]

# --- choices sheet ----------------------------------------------------------
# Every option below was observed in the original paper survey responses. Nothing invented.
CHOICES = [
    ["list_name", "name", "label"],

    ["yes_no", "yes", "Yes"],
    ["yes_no", "no", "No"],

    ["village", "ha_monyake", "Ha Monyake"],
    ["village", "tlokotsing", "Tlokotsing"],
    ["village", "ha_mahlehle", "Ha Mahlehle"],

    ["income_source", "remittance_sa", "Relative working in South Africa"],
    ["income_source", "piece_jobs", "Piece jobs"],
    ["income_source", "child_grant", "Child grants"],
    ["income_source", "pension", "Pension fund"],
    ["income_source", "livestock", "Herding or raising livestock"],
    ["income_source", "crop_sales", "Growing and selling crops"],
    ["income_source", "alcohol_sales", "Making or selling alcohol"],
    ["income_source", "petty_trade", "Selling goods (small trade)"],
    ["income_source", "domestic_work", "Domestic work"],
    ["income_source", "other", "Other"],

    ["strength", "farming", "Farming skills"],
    ["strength", "livestock", "Raising livestock or herding"],
    ["strength", "piece_jobs", "Doing piece jobs"],
    ["strength", "shopkeeping", "Shopkeeping"],
    ["strength", "hairdressing", "Hairdressing"],
    ["strength", "petty_trade", "Selling goods (small trade)"],
    ["strength", "alcohol_production", "Making and selling alcohol"],
    ["strength", "domestic_work", "Domestic work"],
    ["strength", "other", "Other"],

    ["challenge", "adverse_weather", "Adverse weather affecting crops"],
    ["challenge", "land_use_conflict", "Land use conflicts (livestock grazing on farmland)"],
    ["challenge", "food_insecurity", "Lack of food"],
    ["challenge", "no_electricity", "Lack of electricity"],
    ["challenge", "water_shortage", "Water shortages"],
    ["challenge", "poor_roads", "Poor roads"],
    ["challenge", "unemployment", "Struggles to find work"],
    ["challenge", "school_fees", "Cannot afford school fees"],
    ["challenge", "non_payment", "Customers not paying"],
    ["challenge", "seed_supply", "Poor supply of seeds"],
    ["challenge", "weak_collaboration", "Not enough community collaboration"],
    ["challenge", "none", "No challenges reported"],
    ["challenge", "other", "Other"],

    ["opportunity", "electricity", "Supply electricity"],
    ["opportunity", "roads", "Improve roads"],
    ["opportunity", "water_supply", "Better water supply"],
    ["opportunity", "irrigation", "Irrigation (Fato-Fato)"],
    ["opportunity", "farming_inputs", "Distribute farming tools and seeds"],
    ["opportunity", "livestock_support", "Support for raising livestock"],
    ["opportunity", "income_generation", "Create income opportunities"],
    ["opportunity", "community_collaboration", "More community collaboration"],
    ["opportunity", "other", "Other"],
]

# --- settings sheet ---------------------------------------------------------
SETTINGS = [
    ["form_title", "form_id", "version", "default_language"],
    ["CEPA Household Survey", "cepa_household_survey", "2026080701", "English (en)"],
]


def main():
    wb = Workbook()

    ws = wb.active
    ws.title = "survey"
    for row in SURVEY:
        ws.append(row)

    ws = wb.create_sheet("choices")
    for row in CHOICES:
        ws.append(row)

    ws = wb.create_sheet("settings")
    for row in SETTINGS:
        ws.append(row)

    wb.save(OUT)
    print(f"wrote {OUT}: {len(SURVEY)-1} survey rows, {len(CHOICES)-1} choices")


if __name__ == "__main__":
    main()
